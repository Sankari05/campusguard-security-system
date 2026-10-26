"""
generate_backend_documentation.py
---------------------------------
Full generator script that creates:
1. RIDGEWATCH_BACKEND_ARCHITECTURE_DOCUMENTATION.docx
2. RIDGEWATCH_BACKEND_ARCHITECTURE_DOCUMENTATION.pdf

Covers every single backend file, all 25 sections, 6W analysis,
data models, Bayesian fusion engine, domain intelligence, and SQLite persistence.
"""

import os
import sys
from datetime import datetime
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

from generate_backend_documentation_data import BACKEND_FILES_DATA, BACKEND_SECTIONS_CONTENT

print(f"Loaded {len(BACKEND_FILES_DATA)} backend file definitions.")

# ---------------------------------------------------------------------------
# COLOR PALETTE CONSTANTS
# ---------------------------------------------------------------------------
COLOR_PRIMARY_HEX = "0F2547"      # Deep Navy
COLOR_SECONDARY_HEX = "2563EB"    # Accent Ice Blue
COLOR_MUTED_HEX = "64748B"        # Slate Gray
COLOR_BG_LIGHT_HEX = "F3F7FD"     # Light Panel Blue
COLOR_BORDER_HEX = "D3E1F5"       # Border Ice Blue
COLOR_WHITE_HEX = "FFFFFF"

COLOR_PRIMARY = RGBColor(15, 37, 71)
COLOR_SECONDARY = RGBColor(37, 99, 235)
COLOR_MUTED = RGBColor(100, 116, 139)
COLOR_TEXT = RGBColor(15, 23, 42)

RL_PRIMARY = colors.HexColor("#0F2547")
RL_SECONDARY = colors.HexColor("#2563EB")
RL_MUTED = colors.HexColor("#64748B")
RL_BG_LIGHT = colors.HexColor("#F3F7FD")
RL_BORDER = colors.HexColor("#D3E1F5")
RL_TEXT = colors.HexColor("#0F172A")


# ===========================================================================
# DOCX BUILDER
# ===========================================================================

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def add_custom_heading(doc, text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(15)
        run.font.color.rgb = COLOR_PRIMARY
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{COLOR_SECONDARY_HEX}"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(12.5)
        run.font.color.rgb = COLOR_SECONDARY
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_PRIMARY
    return p


def add_paragraph_styled(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = COLOR_PRIMARY
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_TEXT
    return p


def add_callout(doc, text, title="KEY ARCHITECTURAL HIGHLIGHT"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.rows[0].cells[0]
    set_cell_background(cell, COLOR_BG_LIGHT_HEX)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{COLOR_SECONDARY_HEX}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(9)
    r_title.font.color.rgb = COLOR_SECONDARY
    
    r_body = p.add_run(text)
    r_body.font.name = "Arial"
    r_body.font.size = Pt(9)
    r_body.font.color.rgb = COLOR_TEXT
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_row = table.rows[0]
    for i, title in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
        if col_widths and i < len(col_widths):
            cell.width = Inches(col_widths[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
    
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_hex = COLOR_WHITE_HEX if r_idx % 2 == 0 else COLOR_BG_LIGHT_HEX
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_hex)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(str(val))
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.color.rgb = COLOR_TEXT
            
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/></w:tcBorders>')
            tcPr.append(borders)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_file_card_docx(doc, file_info):
    add_custom_heading(doc, f"Module Specification: {file_info['name']}", level=3)
    
    table = doc.add_table(rows=16, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(4.7)
    
    fields = [
        ("FILE NAME", file_info["name"]),
        ("FILE TYPE", file_info["type"]),
        ("PRIMARY PURPOSE", file_info["purpose"]),
        ("WHAT", file_info["what"]),
        ("WHY", file_info["why"]),
        ("WHO", file_info["who"]),
        ("WHERE", file_info["where"]),
        ("WHEN", file_info["when"]),
        ("HOW", file_info["how"]),
        ("INPUTS", file_info["inputs"]),
        ("OUTPUTS", file_info["outputs"]),
        ("DEPENDENCIES", file_info["dependencies"]),
        ("DATA FLOW", file_info["data_flow"]),
        ("DATABASE / PERSISTENCE", file_info["persistence"]),
        ("ERROR HANDLING", file_info["error_handling"]),
        ("IF THIS FILE FAILS", file_info["if_missing"]),
    ]
    
    for idx, (label, val) in enumerate(fields):
        row = table.rows[idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        
        set_cell_background(cell_lbl, COLOR_BG_LIGHT_HEX)
        set_cell_background(cell_val, COLOR_WHITE_HEX)
        set_cell_margins(cell_lbl, top=50, bottom=50, left=70, right=70)
        set_cell_margins(cell_val, top=50, bottom=50, left=70, right=70)
        
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(0)
        p_lbl.paragraph_format.space_after = Pt(0)
        r_lbl = p_lbl.add_run(label)
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(8)
        r_lbl.font.color.rgb = COLOR_PRIMARY
        
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(0)
        p_val.paragraph_format.space_after = Pt(0)
        p_val.paragraph_format.line_spacing = 1.1
        r_val = p_val.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(8)
        r_val.font.color.rgb = COLOR_TEXT
        
        for c in (cell_lbl, cell_val):
            tcPr = c._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/></w:tcBorders>')
            tcPr.append(borders)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build_docx_documentation():
    print("Building Backend DOCX document...")
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # --- TITLE PAGE ---
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pre = p_pre.add_run("RIDGEWATCH RESCUE INTELLIGENCE PLATFORM\n")
    r_pre.font.name = "Arial"
    r_pre.font.size = Pt(11)
    r_pre.font.color.rgb = COLOR_SECONDARY
    r_pre.bold = True
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BACKEND ARCHITECTURE & FILE-BY-FILE CONCEPTUAL SPECIFICATION")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = COLOR_PRIMARY
    r_title.bold = True
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI-Based Multi-Disaster Survivor Detection & Rescue Coordination System\nComprehensive Technical Documentation for Project Review, Mentors, Judges & Maintenance")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = COLOR_MUTED
    
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    
    meta_headers = ["DOCUMENT METRIC", "SYSTEM VALUE", "AUDITED SCOPE"]
    meta_data = [
        ["System Version", "RidgeWatch Backend v2.0 (Production)", "Entire Python Backend Subsystem"],
        ["Execution Runtime", "Python 3.10+ / FastAPI / Uvicorn ASGI Server", "Asynchronous High-Concurrency"],
        ["Total Backend Modules", f"{len(BACKEND_FILES_DATA)} Verified Modules & Datasets", "14 Files in active backend/"],
        ["Database Architecture", "Embedded SQLite with WAL Mode (survivor_detection.db)", "Thread-Safe Persistent Audit Ledger"],
        ["AI Evidence Fusion", "Step 0–Step 9 Bayesian Probabilistic Pipeline", "Explainable SPS Scoring (0–100%)"],
        ["Disaster Taxonomies", "19 Multi-Disaster Domain Intelligence Models", "Landslide, Avalanche, Flood, Earthquake, etc."],
        ["REST Endpoints", "60+ Operational REST & Telemetry Endpoints", "Multi-Operation Lifecycle API"],
        ["Mathematical Integrity", "Strict: Zero Fake Data / Physical Sensor Bounds", "Audited Simulation Physics Models"],
        ["Date of Generation", datetime.now().strftime("%B %d, %Y"), "Production Build Verification"],
    ]
    add_styled_table(doc, meta_headers, meta_data, [2.0, 2.7, 1.8])
    
    add_callout(doc, 
        "This document provides exhaustive, conceptual, file-by-file technical documentation for the complete backend architecture of RidgeWatch. "
        "It details the mathematical Bayesian fusion pipeline, domain hazard intelligence, physics telemetry generation, SQLite persistence, and REST API topology.",
        title="DOCUMENT PURPOSE & EXECUTIVE SCOPE"
    )
    
    doc.add_page_break()
    
    # --- TABLE OF CONTENTS ---
    add_custom_heading(doc, "TABLE OF CONTENTS", level=1)
    toc_items = [
        "1. Executive Summary & Backend Purpose",
        "2. Backend Architectural Philosophy & Design Principles",
        "3. Backend System Architecture & Component Topology",
        "4. Data Models, Schemas & Entity Specifications (models.py)",
        "5. Multi-Sensor Bayesian Fusion & SPS Engine (fusion_engine.py)",
        "6. Scenario Domain Intelligence & Multi-Hazard Rule Engine (scenario_intelligence.py)",
        "7. Sensor & Environmental Physics Simulator (sensor_simulator.py)",
        "8. Autonomous Rescue Dispatch & Staff Eligibility Matching (rescue_dispatch.py)",
        "9. Emergency Alert Generation & Multi-Channel Broadcast (alert_dispatch.py)",
        "10. Hardware Station Registry & Telemetry Diagnostics (hardware_registry.py)",
        "11. Database Architecture, Persistence & Audit Ledger (db.py)",
        "12. FastAPI Application Gateway & REST Endpoints Specification (api.py)",
        "13. Scenario Loader, Master Catalog & Sensor Profiles (scenario_loader.py, scenarios_master.py, scenario_sensor_profiles.py)",
        "14. Standard Scenario Seed Library & Database Scripts (scenarios/*.json, cleanup_default_data.py)",
        "15. Backend-to-Frontend Communication & Telemetry Interfaces",
        "16. End-to-End Backend Operational Workflows (Lifecycle 1 through 10)",
        "17. Backend Data Flow & Processing Pipelines",
        "18. Comprehensive 6W Summary Matrix (All 14 Backend Modules)",
        "19. Standardized Backend File Documentation Blocks (All 14 Files)",
        "20. Component Criticality & Architectural Tier Analysis",
        "21. 'If This File Fails' Systemic Backend Impact Analysis",
        "22. Mathematical & Probabilistic Rigor (SPS, Bayesian, Hazard Multipliers)",
        "23. Security, Validation, Error Handling & Graceful Degradation",
        "24. Backend Architecture Rationale: Why This Backend Architecture?",
        "25. Backend Architecture Summary & Final Conclusions",
    ]
    for item in toc_items:
        p_toc = doc.add_paragraph()
        p_toc.paragraph_format.space_after = Pt(3)
        r_toc = p_toc.add_run(item)
        r_toc.font.name = "Arial"
        r_toc.font.size = Pt(9.5)
        r_toc.font.color.rgb = COLOR_PRIMARY
        r_toc.bold = True
        
    doc.add_page_break()
    
    # --- SECTIONS 1 TO 18 ---
    for sec in BACKEND_SECTIONS_CONTENT[:18]:
        add_custom_heading(doc, sec["title"], level=1)
        for block in sec["blocks"]:
            if block["type"] == "heading2":
                add_custom_heading(doc, block["text"], level=2)
            elif block["type"] == "heading3":
                add_custom_heading(doc, block["text"], level=3)
            elif block["type"] == "paragraph":
                add_paragraph_styled(doc, block["text"], bold_prefix=block.get("prefix"))
            elif block["type"] == "callout":
                add_callout(doc, block["text"], title=block.get("title", "ARCHITECTURAL NOTE"))
            elif block["type"] == "table":
                add_styled_table(doc, block["headers"], block["data"], block.get("widths"))
                
    doc.add_page_break()
    
    # --- SECTION 19: FILE-BY-FILE SPECIFICATION BLOCKS ---
    add_custom_heading(doc, "19. Standardized Backend File Documentation Blocks (All 14 Files)", level=1)
    add_paragraph_styled(doc, "This section provides exhaustive, standardized conceptual specification blocks for all 14 audited files in the backend directory.")
    
    for f_info in BACKEND_FILES_DATA:
        add_file_card_docx(doc, f_info)
        
    doc.add_page_break()
    
    # --- SECTIONS 20 TO 25 ---
    for sec in BACKEND_SECTIONS_CONTENT[19:]:
        add_custom_heading(doc, sec["title"], level=1)
        for block in sec["blocks"]:
            if block["type"] == "heading2":
                add_custom_heading(doc, block["text"], level=2)
            elif block["type"] == "heading3":
                add_custom_heading(doc, block["text"], level=3)
            elif block["type"] == "paragraph":
                add_paragraph_styled(doc, block["text"], bold_prefix=block.get("prefix"))
            elif block["type"] == "callout":
                add_callout(doc, block["text"], title=block.get("title", "ARCHITECTURAL NOTE"))
            elif block["type"] == "table":
                add_styled_table(doc, block["headers"], block["data"], block.get("widths"))
                
    docx_path = "RIDGEWATCH_BACKEND_ARCHITECTURE_DOCUMENTATION.docx"
    doc.save(docx_path)
    print(f"Backend DOCX created successfully at: {docx_path}")
    return docx_path


# ===========================================================================
# PDF BUILDER USING REPORTLAB
# ===========================================================================

class BackendNumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(BackendNumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(BackendNumberedCanvas, self).showPage()
        super(BackendNumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            return  # Skip cover page
        
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(RL_MUTED)
        
        # Header
        self.drawString(54, 750, "RIDGEWATCH — Backend Architecture & File-by-File Conceptual Documentation")
        self.setStrokeColor(RL_BORDER)
        self.setLineWidth(0.5)
        self.line(54, 742, 558, 742)
        
        # Footer
        self.line(54, 45, 558, 45)
        self.drawString(54, 32, "Confidential & Proprietary — RidgeWatch AI Rescue Intelligence Core")
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 32, page_text)
        self.restoreState()


def build_pdf_documentation():
    print("Building Backend PDF document...")
    pdf_path = "RIDGEWATCH_BACKEND_ARCHITECTURE_DOCUMENTATION.pdf"
    
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    style_cover_pre = ParagraphStyle('CoverPre', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, textColor=RL_SECONDARY, alignment=1, spaceAfter=8)
    style_cover_title = ParagraphStyle('CoverTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=RL_PRIMARY, alignment=1, spaceAfter=10)
    style_cover_sub = ParagraphStyle('CoverSub', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14, textColor=RL_MUTED, alignment=1, spaceAfter=20)
    
    style_h1 = ParagraphStyle('DocH1', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=14, leading=17, textColor=RL_PRIMARY, spaceBefore=14, spaceAfter=6, keepWithNext=True)
    style_h2 = ParagraphStyle('DocH2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11.5, leading=14, textColor=RL_SECONDARY, spaceBefore=10, spaceAfter=4, keepWithNext=True)
    style_h3 = ParagraphStyle('DocH3', parent=styles['Heading3'], fontName='Helvetica-Bold', fontSize=10, leading=12, textColor=RL_PRIMARY, spaceBefore=8, spaceAfter=3, keepWithNext=True)
    
    style_body = ParagraphStyle('DocBody', parent=styles['Normal'], fontName='Helvetica', fontSize=8.5, leading=11, textColor=RL_TEXT, spaceAfter=5)
    style_callout = ParagraphStyle('DocCallout', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=10.5, textColor=RL_TEXT)
    style_table_cell = ParagraphStyle('DocTableCell', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=RL_TEXT)
    style_table_hdr = ParagraphStyle('DocTableHdr', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white)
    
    story = []
    
    # --- COVER PAGE ---
    story.append(Spacer(1, 20))
    story.append(Paragraph("RIDGEWATCH RESCUE INTELLIGENCE PLATFORM", style_cover_pre))
    story.append(Paragraph("BACKEND ARCHITECTURE &amp; FILE-BY-FILE CONCEPTUAL SPECIFICATION", style_cover_title))
    story.append(Paragraph("AI-Based Multi-Disaster Survivor Detection &amp; Rescue Coordination System<br/>Comprehensive Technical Documentation for Project Review, Mentors, Judges &amp; Maintenance", style_cover_sub))
    story.append(HRFlowable(width="100%", thickness=1.5, color=RL_SECONDARY, spaceAfter=15))
    
    meta_table_data = [
        [Paragraph("<b>DOCUMENT METRIC</b>", style_table_hdr), Paragraph("<b>SYSTEM VALUE</b>", style_table_hdr), Paragraph("<b>AUDITED SCOPE</b>", style_table_hdr)],
        [Paragraph("System Version", style_table_cell), Paragraph("RidgeWatch Backend v2.0 (Production)", style_table_cell), Paragraph("Entire Python Backend Subsystem", style_table_cell)],
        [Paragraph("Execution Runtime", style_table_cell), Paragraph("Python 3.10+ / FastAPI / Uvicorn ASGI Server", style_table_cell), Paragraph("Asynchronous High-Concurrency", style_table_cell)],
        [Paragraph("Total Backend Modules", style_table_cell), Paragraph(f"{len(BACKEND_FILES_DATA)} Verified Modules &amp; Datasets", style_table_cell), Paragraph("14 Files in active backend/", style_table_cell)],
        [Paragraph("Database Architecture", style_table_cell), Paragraph("Embedded SQLite with WAL Mode (survivor_detection.db)", style_table_cell), Paragraph("Thread-Safe Persistent Audit Ledger", style_table_cell)],
        [Paragraph("AI Evidence Fusion", style_table_cell), Paragraph("Step 0–Step 9 Bayesian Probabilistic Pipeline", style_table_cell), Paragraph("Explainable SPS Scoring (0–100%)", style_table_cell)],
        [Paragraph("Disaster Taxonomies", style_table_cell), Paragraph("19 Multi-Disaster Domain Intelligence Models", style_table_cell), Paragraph("Landslide, Avalanche, Flood, Earthquake, etc.", style_table_cell)],
        [Paragraph("REST Endpoints", style_table_cell), Paragraph("60+ Operational REST &amp; Telemetry Endpoints", style_table_cell), Paragraph("Multi-Operation Lifecycle API", style_table_cell)],
        [Paragraph("Date of Generation", style_table_cell), Paragraph(datetime.now().strftime("%B %d, %Y"), style_table_cell), Paragraph("Production Build Verification", style_table_cell)],
    ]
    t_meta = Table(meta_table_data, colWidths=[1.8*inch, 2.7*inch, 1.9*inch])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), RL_PRIMARY),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, RL_BG_LIGHT]),
        ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 15))
    
    callout_cover = Table([[
        Paragraph("<b>[DOCUMENT PURPOSE &amp; EXECUTIVE SCOPE]</b> This document provides exhaustive, conceptual, file-by-file technical documentation for the complete backend architecture of RidgeWatch. It details the mathematical Bayesian fusion pipeline, domain hazard intelligence, physics telemetry generation, SQLite persistence, and REST API topology.", style_callout)
    ]], colWidths=[6.4*inch])
    callout_cover.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), RL_BG_LIGHT),
        ('LINELEFT', (0, 0), (0, -1), 3.0, RL_SECONDARY),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(callout_cover)
    story.append(PageBreak())
    
    # --- TABLE OF CONTENTS ---
    story.append(Paragraph("TABLE OF CONTENTS", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=RL_SECONDARY, spaceAfter=8))
    toc_items = [
        "1. Executive Summary &amp; Backend Purpose",
        "2. Backend Architectural Philosophy &amp; Design Principles",
        "3. Backend System Architecture &amp; Component Topology",
        "4. Data Models, Schemas &amp; Entity Specifications (models.py)",
        "5. Multi-Sensor Bayesian Fusion &amp; SPS Engine (fusion_engine.py)",
        "6. Scenario Domain Intelligence &amp; Multi-Hazard Rule Engine (scenario_intelligence.py)",
        "7. Sensor &amp; Environmental Physics Simulator (sensor_simulator.py)",
        "8. Autonomous Rescue Dispatch &amp; Staff Eligibility Matching (rescue_dispatch.py)",
        "9. Emergency Alert Generation &amp; Multi-Channel Broadcast (alert_dispatch.py)",
        "10. Hardware Station Registry &amp; Telemetry Diagnostics (hardware_registry.py)",
        "11. Database Architecture, Persistence &amp; Audit Ledger (db.py)",
        "12. FastAPI Application Gateway &amp; REST Endpoints Specification (api.py)",
        "13. Scenario Loader, Master Catalog &amp; Sensor Profiles (scenario_loader.py, scenarios_master.py, scenario_sensor_profiles.py)",
        "14. Standard Scenario Seed Library &amp; Database Scripts (scenarios/*.json, cleanup_default_data.py)",
        "15. Backend-to-Frontend Communication &amp; Telemetry Interfaces",
        "16. End-to-End Backend Operational Workflows (Lifecycle 1 through 10)",
        "17. Backend Data Flow &amp; Processing Pipelines",
        "18. Comprehensive 6W Summary Matrix (All 14 Backend Modules)",
        "19. Standardized Backend File Documentation Blocks (All 14 Files)",
        "20. Component Criticality &amp; Architectural Tier Analysis",
        "21. 'If This File Fails' Systemic Backend Impact Analysis",
        "22. Mathematical &amp; Probabilistic Rigor (SPS, Bayesian, Hazard Multipliers)",
        "23. Security, Validation, Error Handling &amp; Graceful Degradation",
        "24. Backend Architecture Rationale: Why This Backend Architecture?",
        "25. Backend Architecture Summary &amp; Final Conclusions",
    ]
    for item in toc_items:
        story.append(Paragraph(f"<b>{item}</b>", style_body))
    story.append(PageBreak())
    
    # --- SECTIONS 1 TO 18 ---
    for sec in BACKEND_SECTIONS_CONTENT[:18]:
        story.append(Paragraph(sec["title"], style_h1))
        story.append(HRFlowable(width="100%", thickness=0.8, color=RL_SECONDARY, spaceAfter=6))
        for block in sec["blocks"]:
            if block["type"] == "heading2":
                story.append(Paragraph(block["text"], style_h2))
            elif block["type"] == "heading3":
                story.append(Paragraph(block["text"], style_h3))
            elif block["type"] == "paragraph":
                pre = f"<b>{block['prefix']}</b> " if block.get("prefix") else ""
                story.append(Paragraph(f"{pre}{block['text']}", style_body))
            elif block["type"] == "callout":
                t_call = Table([[
                    Paragraph(f"<b>[{block.get('title', 'ARCHITECTURAL NOTE')}]</b> {block['text']}", style_callout)
                ]], colWidths=[6.4*inch])
                t_call.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, -1), RL_BG_LIGHT),
                    ('LINELEFT', (0, 0), (0, -1), 2.5, RL_SECONDARY),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                    ('LEFTPADDING', (0, 0), (-1, -1), 8),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ]))
                story.append(t_call)
                story.append(Spacer(1, 4))
            elif block["type"] == "table":
                t_rows = [[Paragraph(f"<b>{h}</b>", style_table_hdr) for h in block["headers"]]]
                for r in block["data"]:
                    t_rows.append([Paragraph(str(c), style_table_cell) for c in r])
                widths = [w*inch for w in block.get("widths", [2.0, 4.4])]
                t_elem = Table(t_rows, colWidths=widths)
                t_elem.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), RL_PRIMARY),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, RL_BG_LIGHT]),
                    ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
                ]))
                story.append(t_elem)
                story.append(Spacer(1, 4))
        story.append(PageBreak())
        
    # --- SECTION 19: FILE-BY-FILE SPECIFICATION BLOCKS ---
    story.append(Paragraph("19. Standardized Backend File Documentation Blocks (All 14 Files)", style_h1))
    story.append(HRFlowable(width="100%", thickness=0.8, color=RL_SECONDARY, spaceAfter=6))
    story.append(Paragraph("This section provides exhaustive, standardized conceptual specification blocks for all 14 audited files in the backend directory.", style_body))
    
    for f_info in BACKEND_FILES_DATA:
        card_data = [
            [Paragraph("<b>FILE NAME</b>", style_table_cell), Paragraph(f"<b>{f_info['name']}</b>", style_table_cell)],
            [Paragraph("<b>FILE TYPE</b>", style_table_cell), Paragraph(f_info["type"], style_table_cell)],
            [Paragraph("<b>PRIMARY PURPOSE</b>", style_table_cell), Paragraph(f_info["purpose"], style_table_cell)],
            [Paragraph("<b>WHAT</b>", style_table_cell), Paragraph(f_info["what"], style_table_cell)],
            [Paragraph("<b>WHY</b>", style_table_cell), Paragraph(f_info["why"], style_table_cell)],
            [Paragraph("<b>WHO</b>", style_table_cell), Paragraph(f_info["who"], style_table_cell)],
            [Paragraph("<b>WHERE</b>", style_table_cell), Paragraph(f_info["where"], style_table_cell)],
            [Paragraph("<b>WHEN</b>", style_table_cell), Paragraph(f_info["when"], style_table_cell)],
            [Paragraph("<b>HOW</b>", style_table_cell), Paragraph(f_info["how"], style_table_cell)],
            [Paragraph("<b>INPUTS</b>", style_table_cell), Paragraph(f_info["inputs"], style_table_cell)],
            [Paragraph("<b>OUTPUTS</b>", style_table_cell), Paragraph(f_info["outputs"], style_table_cell)],
            [Paragraph("<b>DEPENDENCIES</b>", style_table_cell), Paragraph(f_info["dependencies"], style_table_cell)],
            [Paragraph("<b>DATA FLOW</b>", style_table_cell), Paragraph(f_info["data_flow"], style_table_cell)],
            [Paragraph("<b>DATABASE / PERSISTENCE</b>", style_table_cell), Paragraph(f_info["persistence"], style_table_cell)],
            [Paragraph("<b>ERROR HANDLING</b>", style_table_cell), Paragraph(f_info["error_handling"], style_table_cell)],
            [Paragraph("<b>IF THIS FILE FAILS</b>", style_table_cell), Paragraph(f_info["if_missing"], style_table_cell)],
        ]
        t_card = Table(card_data, colWidths=[1.6*inch, 4.8*inch])
        t_card.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), RL_BG_LIGHT),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
            ('TOPPADDING', (0, 0), (-1, -1), 2.5),
            ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
        ]))
        story.append(KeepTogether([
            Paragraph(f"<b>Module Specification: {f_info['name']}</b>", style_h3),
            t_card,
            Spacer(1, 4)
        ]))
    story.append(PageBreak())
    
    # --- SECTIONS 20 TO 25 ---
    for sec in BACKEND_SECTIONS_CONTENT[19:]:
        story.append(Paragraph(sec["title"], style_h1))
        story.append(HRFlowable(width="100%", thickness=0.8, color=RL_SECONDARY, spaceAfter=6))
        for block in sec["blocks"]:
            if block["type"] == "heading2":
                story.append(Paragraph(block["text"], style_h2))
            elif block["type"] == "heading3":
                story.append(Paragraph(block["text"], style_h3))
            elif block["type"] == "paragraph":
                pre = f"<b>{block['prefix']}</b> " if block.get("prefix") else ""
                story.append(Paragraph(f"{pre}{block['text']}", style_body))
            elif block["type"] == "callout":
                t_call = Table([[
                    Paragraph(f"<b>[{block.get('title', 'ARCHITECTURAL NOTE')}]</b> {block['text']}", style_callout)
                ]], colWidths=[6.4*inch])
                t_call.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, -1), RL_BG_LIGHT),
                    ('LINELEFT', (0, 0), (0, -1), 2.5, RL_SECONDARY),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                    ('LEFTPADDING', (0, 0), (-1, -1), 8),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ]))
                story.append(t_call)
                story.append(Spacer(1, 4))
            elif block["type"] == "table":
                t_rows = [[Paragraph(f"<b>{h}</b>", style_table_hdr) for h in block["headers"]]]
                for r in block["data"]:
                    t_rows.append([Paragraph(str(c), style_table_cell) for c in r])
                widths = [w*inch for w in block.get("widths", [2.0, 4.4])]
                t_elem = Table(t_rows, colWidths=widths)
                t_elem.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), RL_PRIMARY),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ('TOPPADDING', (0, 0), (-1, -1), 3),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, RL_BG_LIGHT]),
                    ('GRID', (0, 0), (-1, -1), 0.5, RL_BORDER),
                ]))
                story.append(t_elem)
                story.append(Spacer(1, 4))
        story.append(PageBreak())
        
    doc.build(story, canvasmaker=BackendNumberedCanvas)
    print(f"Backend PDF created successfully at: {pdf_path}")
    return pdf_path


if __name__ == "__main__":
    docx_file = build_docx_documentation()
    pdf_file = build_pdf_documentation()
    print("Both backend documents generated successfully!")
